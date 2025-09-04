from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
import openpyxl
from catalog.models import Product, Category
from django.utils.text import slugify


BOOLEAN_TRUE = {"true", "1", "yes", "y", "t"}


def parse_bool(value):
    if value is None:
        return False
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in BOOLEAN_TRUE


class Command(BaseCommand):
    help = "Import products from Excel. Creates/updates by SKU if provided else by name."

    def add_arguments(self, parser):
        parser.add_argument("input", help="Path to the Excel file")
        parser.add_argument("--sheet", default="Products", help="Worksheet name")
        parser.add_argument("--dry-run", action="store_true", help="Validate without saving changes")

    @transaction.atomic
    def handle(self, *args, **options):
        path = options["input"]
        sheet_name = options["sheet"]
        dry = options["dry_run"]

        try:
            wb = openpyxl.load_workbook(path)
        except Exception as e:
            raise CommandError(f"Failed to open workbook: {e}")

        if sheet_name not in wb.sheetnames:
            raise CommandError(f"Worksheet '{sheet_name}' not found. Available: {wb.sheetnames}")

        ws = wb[sheet_name]
        headers = [str(ws.cell(row=1, column=c).value or "").strip() for c in range(1, ws.max_column + 1)]

        header_to_idx = {h: i + 1 for i, h in enumerate(headers)}
        required = ["name", "description", "price", "category_name"]
        missing = [r for r in required if r not in header_to_idx]
        if missing:
            raise CommandError(f"Missing required columns: {missing}")

        created, updated, errors = 0, 0, 0
        for r in range(2, ws.max_row + 1):
            row = {h: ws.cell(row=r, column=header_to_idx[h]).value for h in header_to_idx}
            try:
                name = (row.get("name") or "").strip()
                description = (row.get("description") or "").strip()
                price = row.get("price")
                sku = (row.get("sku") or "").strip() or None
                category_name = (row.get("category_name") or "").strip()
                category_slug = (row.get("category_slug") or "").strip()

                if not name or price is None or not category_name:
                    raise ValueError("name, price, category_name are required")

                # Category resolution or creation
                category = None
                if category_slug:
                    category = Category.objects.filter(slug=category_slug).first()
                if not category and category_name:
                    category, _ = Category.objects.get_or_create(name=category_name, defaults={"slug": slugify(category_name)})

                defaults = {
                    "name": name,
                    "description": description,
                    "category": category,
                    "price": price,
                    "stock_quantity": int(row.get("stock_quantity") or 0),
                    "min_stock_level": int(row.get("min_stock_level") or 5),
                    "sku": sku,
                    "weight": row.get("weight_kg") or None,
                    "dimensions": (row.get("dimensions_cm") or "").strip(),
                    "is_active": parse_bool(row.get("is_active")),
                    "is_featured": parse_bool(row.get("is_featured")),
                    "is_available": parse_bool(row.get("is_available", True)),
                    "meta_title": (row.get("meta_title") or "").strip(),
                    "meta_description": (row.get("meta_description") or "").strip(),
                }

                if sku:
                    obj, created_flag = Product.objects.update_or_create(sku=sku, defaults=defaults)
                else:
                    obj, created_flag = Product.objects.update_or_create(name=name, defaults=defaults)

                if created_flag:
                    created += 1
                else:
                    updated += 1

            except Exception as e:
                errors += 1
                self.stderr.write(f"Row {r}: {e}")

        if dry:
            self.stdout.write(self.style.WARNING("Dry run enabled - rolling back changes"))
            raise CommandError(f"Dry run complete. Would create: {created}, update: {updated}, errors: {errors}")

        self.stdout.write(self.style.SUCCESS(f"Imported successfully. Created: {created}, Updated: {updated}, Errors: {errors}"))


