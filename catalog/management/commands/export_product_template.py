from django.core.management.base import BaseCommand
import openpyxl
from openpyxl.utils import get_column_letter


HEADERS = [
    "sku",
    "name",
    "description",
    "category_name",
    "category_slug",
    "price",
    "stock_quantity",
    "min_stock_level",
    "weight_kg",
    "dimensions_cm",
    "is_active",
    "is_featured",
    "is_available",
    "meta_title",
    "meta_description",
]


class Command(BaseCommand):
    help = "Export an Excel template for bulk product import"

    def add_arguments(self, parser):
        parser.add_argument("output", nargs="?", default="product_import_template.xlsx", help="Output file path")

    def handle(self, *args, **options):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Products"

        # Write headers
        for idx, h in enumerate(HEADERS, start=1):
            cell = ws.cell(row=1, column=idx, value=h)
            cell.font = openpyxl.styles.Font(bold=True)
            ws.column_dimensions[get_column_letter(idx)].width = max(14, len(h) + 2)

        # Add notes sheet
        notes = wb.create_sheet(title="Notes")
        notes["A1"] = (
            "Guidelines:\n"
            "- sku: unique identifier (leave blank to auto-generate if supported).\n"
            "- category_name OR category_slug must match an existing category; if not found, a category will be created using category_name.\n"
            "- price: decimal, e.g., 1999.99.\n"
            "- is_active/is_featured/is_available: TRUE/FALSE.\n"
            "- dimensions_cm format example: 10 x 5 x 3.\n"
        )
        notes.merge_cells("A1:D10")

        wb.save(options["output"])
        self.stdout.write(self.style.SUCCESS(f"Template exported to {options['output']}"))


