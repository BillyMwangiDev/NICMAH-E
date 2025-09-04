"""
Normalize varied Excel sheets into the standard product import template.

Input: any .xlsx with columns that can include synonyms like:
- name: ["name", "product name", "DESCRIPTION", "description"]
- sku: ["sku", "SKU", "code"]
- category: ["category", "category_name", "Category", "DESCRIPTION"]
- price: ["price", "unit price", "Price", "Unit Price"]
- stock_quantity: ["stock", "quantity", "current stock", "QUANTITY"]

Output: an Excel file with the headers required by export_product_template.py
on sheet "Products". Fully compatible with import_products.

Usage:
python manage.py normalize_excel_to_template input.xlsx output.xlsx [--sheet Sheet1]
"""

from django.core.management.base import BaseCommand, CommandError
import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter


TEMPLATE_HEADERS = [
    "sku",
    "name",
    "description",
    "category_name",
    "category_slug",
    "price",
    "stock_quantity",
    "min_stock_level",
    "weight_kg",
    "dimensions_cm",
    "is_active",
    "is_featured",
    "is_available",
    "meta_title",
    "meta_description",
]


SYNONYMS = {
    "name": ["name", "product name", "product_name", "DESCRIPTION", "description", "Unnamed: 3"],
    "sku": ["sku", "SKU", "code", "item code"],
    "category_name": ["category", "category_name", "Category", "group", "DESCRIPTION"],
    "price": ["price", "unit price", "Unit Price", "selling price"],
    "stock_quantity": ["stock", "quantity", "current stock", "Current Stock", "QUANTITY", "qty"],
}


def first_matching_col(columns, candidates):
    lower_to_actual = {str(c).strip().lower(): c for c in columns}
    for cand in candidates:
        key = str(cand).strip().lower()
        if key in lower_to_actual:
            return lower_to_actual[key]
    return None


class Command(BaseCommand):
    help = "Normalize an Excel sheet into the standard product import template"

    def add_arguments(self, parser):
        parser.add_argument("input", help="Path to source Excel file")
        parser.add_argument("output", help="Path to write normalized Excel file")
        parser.add_argument("--sheet", default=None, help="Worksheet name (optional)")

    def handle(self, *args, **options):
        input_path = options["input"]
        output_path = options["output"]
        sheet = options["sheet"]

        try:
            df = pd.read_excel(input_path, sheet_name=sheet or 0)
        except FileNotFoundError:
            raise CommandError(f"Input Excel file not found: {input_path}")
        except Exception as exc:
            raise CommandError(f"Failed reading Excel: {exc}")

        # Heuristic: some files have name split across DESCRIPTION + Unnamed: 3
        df.columns = [str(c).strip() for c in df.columns]

        # Build name from best available columns
        name_col = first_matching_col(df.columns, SYNONYMS["name"])  # may be DESCRIPTION or Unnamed: 3
        if name_col in ("DESCRIPTION", "description") and "Unnamed: 3" in df.columns:
            name_series = (df["DESCRIPTION"].fillna("").astype(str) + " " + df["Unnamed: 3"].fillna("").astype(str)).str.strip()
        else:
            if name_col is None and "DESCRIPTION" in df.columns and "Unnamed: 3" in df.columns:
                name_series = (df["DESCRIPTION"].fillna("").astype(str) + " " + df["Unnamed: 3"].fillna("").astype(str)).str.strip()
            elif name_col is not None:
                name_series = df[name_col].fillna("").astype(str).str.strip()
            else:
                name_series = pd.Series(["" for _ in range(len(df))])

        # Other columns
        sku_col = first_matching_col(df.columns, SYNONYMS["sku"]) or None
        category_col = first_matching_col(df.columns, SYNONYMS["category_name"]) or None
        price_col = first_matching_col(df.columns, SYNONYMS["price"]) or None
        stock_col = first_matching_col(df.columns, SYNONYMS["stock_quantity"]) or None

        # Create DataFrame with proper length
        out = pd.DataFrame(index=range(len(df)))
        for h in TEMPLATE_HEADERS:
            out[h] = None
        out["name"] = name_series
        if sku_col:
            out["sku"] = df[sku_col].astype(str).str.strip().replace({"nan": ""})
        else:
            # Auto-generate SKU if missing
            out["sku"] = [
                (n[:3].upper() if isinstance(n, str) and len(n) >= 3 else str(n).upper()).ljust(3, "X") + str(i + 1).zfill(3)
                for i, n in enumerate(out["name"].fillna("").tolist())
            ]

        out["description"] = out["name"].fillna("")
        if category_col:
            out["category_name"] = df[category_col].fillna("").astype(str).str.strip()
        else:
            # Fallback category from DESCRIPTION if present else Uncategorized
            if "DESCRIPTION" in df.columns:
                out["category_name"] = df["DESCRIPTION"].fillna("").astype(str).str.strip().replace({"": "Uncategorized"})
            else:
                out["category_name"] = "Uncategorized"

        out["category_slug"] = ""
        if price_col:
            out["price"] = pd.to_numeric(df[price_col], errors="coerce").fillna(0).round(2)
        else:
            out["price"] = 0

        if stock_col:
            out["stock_quantity"] = pd.to_numeric(df[stock_col], errors="coerce").fillna(0).astype(int)
        else:
            out["stock_quantity"] = 0

        out["min_stock_level"] = 5
        out["weight_kg"] = None
        out["dimensions_cm"] = ""
        out["is_active"] = True
        out["is_featured"] = False
        out["is_available"] = True
        out["meta_title"] = out["name"].fillna("").str[:60]
        out["meta_description"] = out["description"].fillna("").str[:160]

        # Drop empty name rows
        out = out[out["name"].fillna("") != ""].reset_index(drop=True)

        # Write to Excel in the expected sheet
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Products"

            # Headers
            for idx, h in enumerate(TEMPLATE_HEADERS, start=1):
                cell = ws.cell(row=1, column=idx, value=h)
                cell.font = openpyxl.styles.Font(bold=True)
                ws.column_dimensions[get_column_letter(idx)].width = max(14, len(h) + 2)

            # Rows
            for r_idx, (_, row) in enumerate(out.iterrows(), start=2):
                for c_idx, h in enumerate(TEMPLATE_HEADERS, start=1):
                    ws.cell(row=r_idx, column=c_idx, value=row.get(h))

            wb.save(output_path)
        except Exception as exc:
            raise CommandError(f"Failed writing output Excel: {exc}")

        self.stdout.write(self.style.SUCCESS(f"Normalized file written to {output_path} with {len(out)} rows"))


